from flask import Flask, render_template, jsonify, request, send_file
import requests
import subprocess
import logging
import os
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['DEBUG'] = True

# Configuración de logging
logging.basicConfig(level=logging.DEBUG)

def classify_subdomain(subdomain):
    keywords = {
        "files": ["file", "storage", "docs"],
        "backup": ["backup", "bkp", "archive"],
        "conf": ["config", "setup", "settings"],
        "int": ["internal", "intranet"],
        "password": ["auth", "login", "secure"],
        "auth": ["oauth", "authentication"],
        "acl": ["access", "acl"],
        "log": ["log", "report", "audit"],
        "source": ["dev", "source", "src", "code"],
        "exe": ["app", "exe", "bin"],
        "test": ["test", "devtest", "testing"]
    }
    for key, words in keywords.items():
        if any(word in subdomain for word in words):
            return key
    return "int"

def classify_subdomains(subdomains):
    data = []
    for i, subdomain in enumerate(subdomains):
        asset_type = classify_subdomain(subdomain)
        data.append({
            "ID activo": i + 1,
            "Tipo de activo": asset_type,
            "Nombre del activo": subdomain,
            "Descripción": "Subdominio detectado por Api Subdominios by Graviel clasificado como " + asset_type
        })
    return data

def excel_download(data, file_name="activos_digitales.xlsx"):
    df = pd.DataFrame(data)
    file_path = os.path.join(os.getcwd(), file_name)
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Activos Digitales")
        workbook = writer.book
        worksheet = writer.sheets["Activos Digitales"]
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = length
        table = Table(displayName="TablaActivos", ref=worksheet.dimensions)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        worksheet.add_table(table)
    logging.debug(f"Archivo Excel guardado en: {file_path}")
    return file_path

@app.route('/')
def home():
    return render_template('index.html')

@app.route('/assetfinder', methods=['POST'])
def assetfinder():
    domain = request.form.get('domain')
    if not domain:
        return render_template('result_assetfinder.html', error='Dominio no proporcionado')

    try:
        result = subprocess.run(['assetfinder', domain], capture_output=True, text=True)
        output = result.stdout
        subdomains = output.splitlines()
    except Exception as e:
        logging.error(f"Error al ejecutar assetfinder: {e}")
        return render_template('result_assetfinder.html', error=str(e))

    data = classify_subdomains(subdomains)
    file_path = excel_download(data)

    return render_template('result_assetfinder.html', subdomains=subdomains, file_path=file_path)

@app.route('/subdominios', methods=['POST'])
def subdominios():
    domain = request.form.get('domain')
    if not domain:
        return render_template('resultApiSubdomain.html', error='Dominio no proporcionado')

    try:
        url = "https://subdomain-finder3.p.rapidapi.com/v1/subdomain-finder/"
        querystring = {"domain": domain}
        headers = {
            "x-rapidapi-key": "e1394062c0mshd5dc5bcaab94015p1206adjsn94a668f30de2",
            "x-rapidapi-host": "subdomain-finder3.p.rapidapi.com"
        }
        response = requests.get(url, headers=headers, params=querystring)
        response.raise_for_status()
        response_data = response.json()
        subdomains = response_data.get('subdomains', [])
    except requests.exceptions.RequestException as e:
        logging.error(f"Error al llamar a la API de RapidAPI: {e}")
        return render_template('resultApiSubdomain.html', error=f"Error al llamar a la API de RapidAPI: {e}")
    except Exception as e:
        logging.error(f"Error general: {e}")
        return render_template('resultApiSubdomain.html', error=str(e))

    data = classify_subdomains(subdomains)
    file_path = excel_download(data, file_name="subdominios_api.xlsx")

    return render_template('resultApiSubdomain.html', subdomains=subdomains, file_path=file_path)

@app.route('/whois', methods=['POST'])
def whois():
    domain = request.form.get('domain')
    if not domain:
        return render_template('resultadosWhois.html', error='Dominio no proporcionado')

    try:
        url = "https://www.whoisxmlapi.com/whoisserver/WhoisService"
        params = {
            "apiKey": "at_qT2z1yWbC5G99r0FRcrxhzO2R4e7N",
            "domainName": domain,
            "outputFormat": "JSON"
        }
        response = requests.get(url, params=params)
        response.raise_for_status()
        response_data = response.json()
    except requests.exceptions.RequestException as e:
        logging.error(f"Error al llamar a la API de WHOISXML: {e}")
        return render_template('resultadosWhois.html', error=f"Error al llamar a la API de WHOISXML: {e}")
    except Exception as e:
        logging.error(f"Error general: {e}")
        return render_template('resultadosWhois.html', error=str(e))

    whois_data = response_data
    data = [{"ID activo": i + 1, "Tipo de activo": "whois", "Nombre del activo": key, "Descripción": value} for i, (key, value) in enumerate(whois_data.items())]
    file_path = excel_download(data, file_name="whois.xlsx")

    return render_template('resultadosWhois.html', whois_data=whois_data, file_path=file_path)

@app.route('/amass', methods=['POST'])
def amass():
    domain = request.form.get('domain')
    if not domain:
        return render_template('resultAmass.html', error='Dominio no proporcionado')

    command = f"amass enum -d {domain}"
    try:
        result = subprocess.run(command, shell=True, capture_output=True, text=True)
        output = result.stdout
        if result.returncode != 0:
            raise Exception(f"Error en la ejecución de Amass: {output}")
        subdomains = output.splitlines()
    except Exception as e:
        logging.error(f"Error al ejecutar Amass: {e}")
        return render_template('resultAmass.html', error=f"Error al ejecutar Amass: {e}")

    data = classify_subdomains(subdomains)
    file_path = excel_download(data, file_name="amass.xlsx")

    return render_template('resultAmass.html', subdomains=subdomains, file_path=file_path)

@app.route('/theharvester', methods=['POST'])
def theharvester():
    domain = request.form.get('domain')
    if not domain:
        return render_template('resultHarvester.html', error='Dominio no proporcionado')

    command = f"theHarvester -d {domain} -b all"
    try:
        result = subprocess.run(command, shell=True, capture_output=True, text=True)
        output = result.stdout
        if result.returncode != 0:
            raise Exception(f"Error en la ejecución de The Harvester: {output}")
        emails = [line for line in output.splitlines() if '@' in line]
    except Exception as e:
        logging.error(f"Error al ejecutar The Harvester: {e}")
        return render_template('resultHarvester.html', error=f"Error al ejecutar The Harvester: {e}")

    data = [{"ID activo": i + 1, "Tipo de activo": "email", "Nombre del activo": email, "Descripción": "Email encontrado por The Harvester"} for i, email in enumerate(emails)]
    file_path = excel_download(data, file_name="theharvester.xlsx")

    return render_template('resultHarvester.html', emails=emails, file_path=file_path)

@app.route('/download', methods=['GET'])
def download():
    file_path = request.args.get('file_path')
    logging.debug(f"Intentando descargar el archivo en la ruta: {file_path}")
    if file_path and os.path.exists(file_path):
        return send_file(file_path, as_attachment=True)
    else:
        return jsonify({"message": "File not found"}), 404

if __name__ == "__main__":
    app.run(debug=True)